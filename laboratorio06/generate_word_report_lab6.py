import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_report_lab6():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    COLOR_TECSUP_BLUE = RGBColor(0, 74, 143)
    COLOR_GRAY = RGBColor(108, 117, 125)

    # PORTADA
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_inst = p_inst.add_run("TECSUP — DEPARTAMENTO DE TECNOLOGÍA DIGITAL")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(10)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("LABORATORIO N° 6\nMOTOR DE PLANTILLAS CON DJANGO")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_TECSUP_BLUE

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Desarrollo de Aplicaciones Empresariales — Sección 4-C24-A")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_GRAY

    doc.add_paragraph()
    doc.add_paragraph()

    table_info = doc.add_table(rows=4, cols=2)
    table_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_info.autofit = False

    datos = [
        ("Estudiante:", "Israel Angel Huaman Huaman"),
        ("Docente:", "Michael Montgomery Rosell"),
        ("Fecha de entrega:", "Octubre 2026"),
        ("Entorno de desarrollo:", "Python 3.14 / Django 6.1 / SQLite3 / Edge"),
    ]

    for i, (label, val) in enumerate(datos):
        cell_lbl = table_info.cell(i, 0)
        cell_val = table_info.cell(i, 1)
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        p1 = cell_lbl.paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(11)

        p2 = cell_val.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.name = "Arial"
        r2.font.size = Pt(11)
        if label == "Estudiante:":
            r2.font.bold = True
            r2.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_page_break()

    # 1. CAPACIDADES Y OBJETIVOS
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("1. Capacidades y Objetivos del Laboratorio")
    r1.font.name = "Arial"
    r1.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph(
        "En este laboratorio se implementa un portal de noticias web corporativo en Django, enfocado en el dominio del motor de plantillas de Django (DTL). "
        "Se desarrollan las capacidades formales establecidas en el documento de trabajo:"
    )

    caps = [
        "Implementar el motor de plantillas de Django con herencia (base.html) y fragmentos modulares reutilizables (_article_card.html).",
        "Formatear y condicionar datos en la capa de presentación mediante variables, filtros (|date, |time, |truncatewords) y etiquetas de control ({% for %}, {% empty %}, {% if %}).",
        "Administrar las entidades Article, Category y Author desde el panel de control y publicarlas dinámicamente en las plantillas.",
        "Comprobar y documentar el mecanismo de autoescapado (autoescape) de Django como defensa activa frente a ataques Cross-Site Scripting (XSS)."
    ]
    for c in caps:
        doc.add_paragraph(c, style='List Bullet')

    # 2. ESTRUCTURA DEL PROYECTO
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("2. Estructura del Proyecto en el Editor")
    r2.font.name = "Arial"
    r2.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph("Estructura de archivos y directorios del proyecto en laboratorio06:")
    p_code = doc.add_paragraph()
    r_code = p_code.add_run(
        "laboratorio06/\n"
        "├── config/                    # Configuración global del proyecto\n"
        "├── news/                      # Aplicación de noticias y artículos\n"
        "│   ├── models.py              # Article, Category, Author\n"
        "│   ├── admin.py               # ModelAdmin con search, filter y list_display\n"
        "│   ├── views.py               # Vistas: home, category_articles, article_detail\n"
        "│   ├── urls.py                # Enrutador con nombres propios\n"
        "│   ├── tests.py               # 5 pruebas automáticas (herencia, fragmentos, XSS)\n"
        "│   ├── context_processors.py  # Categorías globales en barra lateral\n"
        "│   └── management/commands/   # setup_admin.py y seed_news.py\n"
        "├── templates/\n"
        "│   ├── base.html              # Plantilla base con bloques: title, content, sidebar\n"
        "│   └── news/\n"
        "│       ├── _article_card.html      # Fragmento modular de tarjeta de noticia\n"
        "│       ├── home.html               # Portada con bucle for, empty y truncatewords\n"
        "│       ├── category_articles.html  # Listado por categoría reutilizando tarjeta\n"
        "│       └── article_detail.html     # Ficha de detalle y prueba de autoescape\n"
        "├── static/css/style.css       # Hoja de estilos vinculada con {% load static %}\n"
        "├── requirements.txt           # Dependencias del proyecto\n"
        "├── README.md                  # Documentación técnica completa\n"
        "└── manage.py                  # CLI de Django"
    )
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(9.5)

    doc.add_page_break()

    # 3. HERENCIA Y FRAGMENTOS REUTILIZABLES
    h3 = doc.add_heading(level=1)
    r3 = h3.add_run("3. Motor de Plantillas: Herencia y Fragmentos Reutilizables (Requisitos 4, 5, 6 y 8)")
    r3.font.name = "Arial"
    r3.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph(
        "Se aplicó el principio DRY (Don't Repeat Yourself) implementando una plantilla base (base.html) con bloques semánticos "
        "y un fragmento reutilizable (_article_card.html) que encapsula el marcado de cada noticia."
    )

    doc.add_heading("3.1. Portada del Portal (home.html)", level=2)
    doc.add_paragraph(
        "La portada recorre los artículos usando {% for %}, incluye _article_card.html para cada elemento y dispone de la cláusula {% empty %} "
        "para el caso de no encontrar resultados. Los resúmenes son recortados a nivel de presentación con |truncatewords:25:"
    )
    if os.path.exists("screenshots/01_portada_noticias.png"):
        doc.add_picture("screenshots/01_portada_noticias.png", width=Inches(6.2))
        p_c = doc.add_paragraph("Figura 1: Portada del portal con tarjetas modulares, barra lateral dinámica y filtros aplicados.")
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c.style.font.size = Pt(9)
        p_c.style.font.italic = True

    doc.add_heading("3.2. Listado por Categoría Reutilizando el Mismo Fragmento (category_articles.html)", level=2)
    doc.add_paragraph(
        "Al ingresar a una categoría (ej. Tecnología), se reutiliza exactamente el mismo componente _article_card.html, "
        "garantizando consistencia visual y mantenimiento centralizado del código HTML:"
    )
    if os.path.exists("screenshots/02_categoria_tecnologia.png"):
        doc.add_picture("screenshots/02_categoria_tecnologia.png", width=Inches(6.2))
        p_c = doc.add_paragraph("Figura 2: Listado filtrado por categoría 'Tecnología' reutilizando _article_card.html.")
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c.style.font.size = Pt(9)
        p_c.style.font.italic = True

    doc.add_page_break()

    # 4. FICHA DE DETALLE Y ESTILOS
    h4 = doc.add_heading(level=1)
    r4 = h4.add_run("4. Ficha de Detalle de Noticia y Archivos Estáticos (Requisitos 7, 9 y 10)")
    r4.font.name = "Arial"
    r4.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph(
        "La vista de detalle hereda de base.html e inyecta la información del autor, categorías asociadas, copete y contenido completo. "
        "Los estilos globales se cargan mediante {% load static %} desde static/css/style.css, y todos los enlaces se construyen usando {% url %}:"
    )
    if os.path.exists("screenshots/03_detalle_noticia.png"):
        doc.add_picture("screenshots/03_detalle_noticia.png", width=Inches(6.2))
        p_c = doc.add_paragraph("Figura 3: Vista de detalle de artículo con autor, categorías, migas de pan y estilos estáticos.")
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c.style.font.size = Pt(9)
        p_c.style.font.italic = True

    # 5. DEMOSTRACIÓN DE ESCAPADO AUTOMÁTICO (AUTOESCAPE)
    h5 = doc.add_heading(level=1)
    r5 = h5.add_run("5. Demostración de Escapado Automático (Autoescape) y Seguridad XSS (Requisito 12)")
    r5.font.name = "Arial"
    r5.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph(
        "Para responder al Requisito 12, se guardó en el cuerpo de una noticia una etiqueta de script maliciosa: "
        "<script>alert('Ataque XSS');</script> y una etiqueta <b>Texto</b>:"
    )
    if os.path.exists("screenshots/04_demostracion_autoescape.png"):
        doc.add_picture("screenshots/04_demostracion_autoescape.png", width=Inches(6.2))
        p_c = doc.add_paragraph("Figura 4: Demostración de autoescape — Las etiquetas HTML no se ejecutan y se muestran como texto seguro.")
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c.style.font.size = Pt(9)
        p_c.style.font.italic = True

    doc.add_paragraph(
        "Explicación Técnica:\n"
        "El motor de plantillas de Django cuenta con Autoescape activado por defecto. Cuando una variable es renderizada, "
        "Django convierte los caracteres peligrosos en sus entidades HTML correspondientes (< se transforma en &lt;, > en &gt;, etc.). "
        "Gracias a esto, el navegador interpreta el script como una simple cadena de texto y no como código ejecutable en el DOM, "
        "bloqueando ataques de inyección Cross-Site Scripting (XSS)."
    )

    doc.add_page_break()

    # 6. GESTIÓN DESDE EL ADMINISTRADOR
    h6 = doc.add_heading(level=1)
    r6 = h6.add_run("6. Gestión de Contenidos desde Django Admin (Requisito 11)")
    r6.font.name = "Arial"
    r6.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph(
        "Se personalizaron las tres entidades (Article, Category, Author) utilizando ModelAdmin con list_display, list_filter, "
        "search_fields, prepopulated_fields y filter_horizontal. Se cargaron 6 noticias en 3 categorías diferentes:"
    )
    if os.path.exists("screenshots/05_admin_articulos.png"):
        doc.add_picture("screenshots/05_admin_articulos.png", width=Inches(6.2))
        p_c = doc.add_paragraph("Figura 5: Gestión de artículos en Django Admin con filtros por estado, fecha y categorías.")
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c.style.font.size = Pt(9)
        p_c.style.font.italic = True

    # 7. PRUEBAS AUTOMATIZADAS
    h7 = doc.add_heading(level=1)
    r7 = h7.add_run("7. Casos de Prueba Automatizados")
    r7.font.name = "Arial"
    r7.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph("Ejecución de la suite de pruebas unitarias en news/tests.py con 'python manage.py test':")
    p_tb = doc.add_paragraph()
    r_tb = p_tb.add_run(
        "Creating test database for alias 'default'...\n"
        "test_article_detail_view (news.tests.TemplateEngineTestCase) ... ok\n"
        "test_autoescape_protection_xss (news.tests.TemplateEngineTestCase) ... ok\n"
        "test_category_view_status_and_reusable_fragment (news.tests.TemplateEngineTestCase) ... ok\n"
        "test_empty_state_in_home (news.tests.TemplateEngineTestCase) ... ok\n"
        "test_home_view_status_and_templates (news.tests.TemplateEngineTestCase) ... ok\n\n"
        "----------------------------------------------------------------------\n"
        "Ran 5 tests in 0.215s\n\n"
        "OK"
    )
    r_tb.font.name = "Consolas"
    r_tb.font.size = Pt(9)

    # 8. CONCLUSIONES
    h8 = doc.add_heading(level=1)
    r8 = h8.add_run("8. Conclusiones")
    r8.font.name = "Arial"
    r8.font.color.rgb = COLOR_TECSUP_BLUE

    conclusiones = [
        "1. La herencia de plantillas en Django permite desacoplar el layout estructural de la lógica de presentación específica de cada vista.",
        "2. El uso de fragmentos reutilizables ({% include %}) evita la duplicación de marcado y facilita el mantenimiento global del diseño del portal.",
        "3. Los filtros de plantilla nativos (|date, |truncatewords) resuelven las necesidades de formateo directamente en la capa de vista sin contaminar los modelos de negocio.",
        "4. El mecanismo de autoescape por defecto de Django es un pilar fundamental de seguridad corporativa, mitigando ataques de Cross-Site Scripting (XSS) sin requerir configuración manual."
    ]
    for con in conclusiones:
        doc.add_paragraph(con)

    output_filename = "Laboratorio6_Israel_Huaman.docx"
    doc.save(output_filename)
    print(f"Informe Word generado con éxito: {output_filename}")

if __name__ == "__main__":
    generate_report_lab6()
