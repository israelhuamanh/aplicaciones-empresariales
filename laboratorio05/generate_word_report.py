import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_report():
    doc = docx.Document()

    # Configurar márgenes de página (1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Colores institucionales
    COLOR_TECSUP_BLUE = RGBColor(0, 74, 143)
    COLOR_DARK = RGBColor(33, 37, 41)
    COLOR_GRAY = RGBColor(108, 117, 125)

    # --- PORTADA ---
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_inst = p_inst.add_run("TECSUP — DEPARTAMENTO DE TECNOLOGÍA DIGITAL")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(10)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph() # Espacio

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("LABORATORIO N° 5\nADMINISTRADOR CON DJANGO")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_TECSUP_BLUE

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Desarrollo de Aplicaciones Empresariales — Sección 4-C24-A")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_GRAY

    doc.add_paragraph() # Espacio
    doc.add_paragraph()

    # Tabla de datos informativos
    table_info = doc.add_table(rows=4, cols=2)
    table_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_info.autofit = False

    datos = [
        ("Estudiante:", "Israel Angel Huaman Huaman"),
        ("Docente:", "Michael Montgomery Rosell"),
        ("Fecha de entrega:", "Octubre 2026"),
        ("Entorno de desarrollo:", "Python 3.14 / Django 6.1 / SQLite3 / Edge"),
    ]

    for i, (label, val) in enumerate(datos):
        cell_lbl = table_info.cell(i, 0)
        cell_val = table_info.cell(i, 1)
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        p1 = cell_lbl.paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(11)

        p2 = cell_val.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.name = "Arial"
        r2.font.size = Pt(11)
        if label == "Estudiante:":
            r2.font.bold = True
            r2.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_page_break()

    # --- SECCIÓN 1: INTRODUCCIÓN Y CAPACIDADES ---
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Capacidades y Objetivos del Laboratorio")
    r_h1.font.name = "Arial"
    r_h1.font.color.rgb = COLOR_TECSUP_BLUE

    p_desc = doc.add_paragraph(
        "El presente laboratorio implementa un sistema empresarial robusto en Django para la gestión cinematográfica (películas, géneros, personas y valoraciones). "
        "En este entregable se demuestran de forma exhaustiva las capacidades solicitadas en la guía oficial:"
    )
    p_desc.style.font.name = "Arial"
    p_desc.style.font.size = Pt(11)

    caps = [
        "Configuración del administrador de Django para gestionar modelos altamente relacionados (Movie, Genre, Person, Rating).",
        "Personalización del listado, filtros por facetas y búsqueda avanzada mediante ModelAdmin.",
        "Integración de bloques de líneas (TabularInline) para editar valoraciones dentro del formulario padre.",
        "Implementación de campos de solo lectura para auditoría estricta (created_at y updated_at).",
        "Control de acceso basado en roles (RBAC) mediante el grupo «editores», validando la restricción de eliminación.",
        "Desarrollo de una vista pública de recomendación con agregaciones ORM y filtros interactivos por género."
    ]
    for c in caps:
        p_item = doc.add_paragraph(c, style='List Bullet')
        p_item.style.font.name = "Arial"
        p_item.style.font.size = Pt(10.5)

    # --- SECCIÓN 2: ESTRUCTURA DEL PROYECTO ---
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Estructura del Proyecto en el Editor")
    r_h2.font.name = "Arial"
    r_h2.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph(
        "El proyecto fue construido desde cero con una arquitectura modular limpia siguiendo la guía oficial y el estándar PEP 8:"
    )

    p_tree = doc.add_paragraph()
    r_tree = p_tree.add_run(
        "laboratorio5/\n"
        "├── config/                    # Configuración central (settings, urls, wsgi)\n"
        "├── movies/                    # Aplicación de películas\n"
        "│   ├── models.py              # Modelos Movie, Genre, Person, Rating\n"
        "│   ├── admin.py               # MovieAdmin, RatingInline, filtros, auditoría\n"
        "│   ├── views.py               # Vista de recomendaciones y detalle\n"
        "│   ├── urls.py                # Enrutador de la app movies\n"
        "│   ├── tests.py               # 10 pruebas unitarias automatizadas\n"
        "│   ├── management/commands/   # seed_data.py y setup_roles.py\n"
        "│   └── templates/movies/      # base.html, recommendations.html, movie_detail.html\n"
        "├── requirements.txt           # Dependencias exactas (Django, Pillow)\n"
        "├── README.md                  # Documentación técnica completa\n"
        "└── manage.py                  # CLI de gestión Django"
    )
    r_tree.font.name = "Consolas"
    r_tree.font.size = Pt(9.5)

    # --- SECCIÓN 3: PERSONALIZACIÓN DEL PANEL DE ADMINISTRACIÓN ---
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Personalización del Panel de Administración (Requisitos 4, 5, 6 y 7)")
    r_h3.font.name = "Arial"
    r_h3.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph(
        "A continuación se presenta el contraste visual del panel administrativo antes y después de la personalización con ModelAdmin, "
        "así como el uso de edición en línea y campos de auditoría protegidos."
    )

    doc.add_heading("3.1. Dashboard Principal del Administrador", level=2)
    doc.add_paragraph("Panel general con los cuatro modelos registrados y gestionables de forma directa:")
    if os.path.exists("screenshots/04_admin_dashboard.png"):
        doc.add_picture("screenshots/04_admin_dashboard.png", width=Inches(6.2))
        p_cap = doc.add_paragraph("Figura 1: Dashboard de Django Admin con los modelos Movie, Genre, Person y Rating.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.style.font.size = Pt(9)
        p_cap.style.font.italic = True

    doc.add_heading("3.2. Listado Avanzado con Filtros y Búsqueda (list_display, list_filter, search_fields)", level=2)
    doc.add_paragraph(
        "Se sustituyó el listado básico por una vista integral que muestra: Título, Año de estreno, Duración, Director, Géneros asociados "
        "y el Promedio de Valoración calculado dinámicamente. A la derecha se disponen los filtros por género y año de estreno, y arriba la barra de búsqueda por título o director:"
    )
    if os.path.exists("screenshots/05_admin_peliculas_personalizado.png"):
        doc.add_picture("screenshots/05_admin_peliculas_personalizado.png", width=Inches(6.2))
        p_cap = doc.add_paragraph("Figura 2: Listado de películas con columnas personalizadas, filtros laterales y búsqueda inteligente.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.style.font.size = Pt(9)
        p_cap.style.font.italic = True

    doc.add_heading("3.3. Edición en Línea (RatingInline) y Campos de Auditoría Protegidos", level=2)
    doc.add_paragraph(
        "Al ingresar a una película (ej. Inception), se observa:\n"
        "1. Campos de Auditoría (created_at y updated_at) como de solo lectura (readonly_fields), impidiendo su manipulación manual.\n"
        "2. Bloque de valoraciones en línea (TabularInline), que permite registrar críticas y notas del 1 al 10 sin salir del formulario padre."
    )
    if os.path.exists("screenshots/06_admin_formulario_inlines_auditoria.png"):
        doc.add_picture("screenshots/06_admin_formulario_inlines_auditoria.png", width=Inches(6.2))
        p_cap = doc.add_paragraph("Figura 3: Formulario padre con bloque tabular de valoraciones y auditoría no editable.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.style.font.size = Pt(9)
        p_cap.style.font.italic = True

    doc.add_page_break()

    # --- SECCIÓN 4: COMPARACIÓN DE ACCESOS Y PERMISOS ---
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Comparación de Accesos: Superusuario vs. Grupo «editores» (Requisito 9 y 11)")
    r_h4.font.name = "Arial"
    r_h4.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph(
        "Para cumplir con el Requisito 9, se creó el grupo «editores» asignándole permisos para añadir (add_movie), "
        "modificar (change_movie) y visualizar (view_movie), pero RESTRINGIENDO ESTRICTAMENTE la eliminación (delete_movie). "
        "A continuación se contrasta visualmente lo que observa el Superusuario (admin) frente al Usuario Editor (editor1):"
    )

    doc.add_heading("4.1. Vista del Formulario para el Superusuario (Acceso Total)", level=2)
    doc.add_paragraph("El superusuario dispone del botón rojo 'Eliminar' en la parte inferior izquierda:")
    if os.path.exists("screenshots/07_admin_botones_superuser.png"):
        doc.add_picture("screenshots/07_admin_botones_superuser.png", width=Inches(6.2))
        p_cap = doc.add_paragraph("Figura 4: Superusuario — Dispone de permiso para eliminar registros.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.style.font.size = Pt(9)
        p_cap.style.font.italic = True

    doc.add_heading("4.2. Vista del Formulario para el Usuario Editor (Restricción de Borrado)", level=2)
    doc.add_paragraph(
        "Al iniciar sesión como 'editor1', el botón 'Eliminar' desaparece por completo del formulario, "
        "y en el listado masivo ya no figura la acción de borrar. El editor únicamente puede guardar y modificar:"
    )
    if os.path.exists("screenshots/10_editor_botones_sin_eliminar.png"):
        doc.add_picture("screenshots/10_editor_botones_sin_eliminar.png", width=Inches(6.2))
        p_cap = doc.add_paragraph("Figura 5: Usuario Editor — El botón 'Eliminar' desaparece y no puede borrar películas.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.style.font.size = Pt(9)
        p_cap.style.font.italic = True

    doc.add_page_break()

    # --- SECCIÓN 5: VISTA PÚBLICA DE RECOMENDACIONES ---
    h5 = doc.add_heading(level=1)
    r_h5 = h5.add_run("5. Vista Pública de Recomendaciones (Requisito 10)")
    r_h5.font.name = "Arial"
    r_h5.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph(
        "Se construyó una interfaz web moderna accesible públicamente en '/movies/', demostrando la diferencia entre "
        "la gestión administrativa privada y el consumo de datos para el usuario final mediante el ORM de Django."
    )

    doc.add_heading("5.1. Cartelera de Películas Mejor Valoradas y Filtros Dinámicos", level=2)
    doc.add_paragraph(
        "La vista calcula el promedio de valoraciones en tiempo real (Avg('ratings__score')) y ordena las películas de mayor a menor puntuación. "
        "Permite además filtrar por género con botones interactivos:"
    )
    if os.path.exists("screenshots/01_recomendaciones_publicas.png"):
        doc.add_picture("screenshots/01_recomendaciones_publicas.png", width=Inches(6.2))
        p_cap = doc.add_paragraph("Figura 6: Vista pública de recomendaciones con promedio de estrellas y tarjetas responsivas.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.style.font.size = Pt(9)
        p_cap.style.font.italic = True

    doc.add_heading("5.2. Filtrado por Género (Ciencia Ficción)", level=2)
    doc.add_paragraph("Al seleccionar un género específico, el catálogo se actualiza automáticamente mostrando los títulos mejor puntuados de dicha categoría:")
    if os.path.exists("screenshots/02_filtro_genero.png"):
        doc.add_picture("screenshots/02_filtro_genero.png", width=Inches(6.2))
        p_cap = doc.add_paragraph("Figura 7: Catálogo filtrado por el género 'Ciencia Ficción'.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.style.font.size = Pt(9)
        p_cap.style.font.italic = True

    doc.add_heading("5.3. Ficha Técnica y Detalle de Críticas", level=2)
    doc.add_paragraph("Cada película enlaza a su vista individual (/movies/<id>/) con sinopsis, directores, reparto y el desglose de valoraciones:")
    if os.path.exists("screenshots/03_detalle_pelicula.png"):
        doc.add_picture("screenshots/03_detalle_pelicula.png", width=Inches(6.2))
        p_cap = doc.add_paragraph("Figura 8: Ficha de detalle de Inception con desglose de críticas y valoraciones.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.style.font.size = Pt(9)
        p_cap.style.font.italic = True

    doc.add_page_break()

    # --- SECCIÓN 6: PRUEBAS AUTOMATIZADAS ---
    h6 = doc.add_heading(level=1)
    r_h6 = h6.add_run("6. Casos de Prueba y Verificación Automatizada")
    r_h6.font.name = "Arial"
    r_h6.font.color.rgb = COLOR_TECSUP_BLUE

    doc.add_paragraph(
        "Para garantizar la calidad de software requerida por un perfil senior, se implementó una batería de 10 pruebas automatizadas en 'movies/tests.py'. "
        "A continuación se presenta el resultado de la ejecución con 'python manage.py test -v 2':"
    )

    test_box = doc.add_paragraph()
    r_tb = test_box.add_run(
        "test_audit_timestamps (movies.tests.MovieModelTestCase) ... ok\n"
        "test_model_str_representations (movies.tests.MovieModelTestCase) ... ok\n"
        "test_movie_relationships_and_aggregations (movies.tests.MovieModelTestCase) ... ok\n"
        "test_editor_cannot_delete_in_admin (movies.tests.PermissionsAndRolesTestCase) ... ok\n"
        "test_group_has_add_and_change_but_not_delete (movies.tests.PermissionsAndRolesTestCase) ... ok\n"
        "test_normal_user_cannot_access_admin (movies.tests.PermissionsAndRolesTestCase) ... ok\n"
        "test_superuser_can_access_delete_in_admin (movies.tests.PermissionsAndRolesTestCase) ... ok\n"
        "test_movie_detail_view (movies.tests.RecommendationsViewTestCase) ... ok\n"
        "test_recommendations_genre_filtering (movies.tests.RecommendationsViewTestCase) ... ok\n"
        "test_recommendations_page_status_and_template (movies.tests.RecommendationsViewTestCase) ... ok\n\n"
        "----------------------------------------------------------------------\n"
        "Ran 10 tests in 10.448s\n\n"
        "OK"
    )
    r_tb.font.name = "Consolas"
    r_tb.font.size = Pt(9)

    doc.add_paragraph(
        "Todos los tests concluyeron exitosamente (OK), validando:\n"
        "• Cálculo aritmético de promedios en Movie.\n"
        "• Restricción HTTP 403 Forbidden para editores al intentar eliminar películas.\n"
        "• Acceso HTTP 200 y permisos completos para el superusuario.\n"
        "• Bloqueo y redirección HTTP 302 a usuarios sin rango staff."
    )

    # --- SECCIÓN 7: CONCLUSIONES ---
    h7 = doc.add_heading(level=1)
    r_h7 = h7.add_run("7. Conclusiones")
    r_h7.font.name = "Arial"
    r_h7.font.color.rgb = COLOR_TECSUP_BLUE

    conclusiones = [
        "1. Django Admin es una herramienta empresarial de alta productividad que permite administrar relaciones ManyToMany y ForeignKey con inlines ergonómicos sin programar vistas adicionales.",
        "2. El sistema de permisos y grupos de Django permite implementar esquemas de seguridad basados en roles (RBAC) con precisión milimétrica, protegiendo operaciones críticas como la eliminación de registros.",
        "3. La auditoría mediante campos de solo lectura (readonly_fields) garantiza la integridad histórica de los registros, evitando adulteraciones manuales.",
        "4. El ORM de Django facilita la agregación y el ordenamiento eficiente de consultas complejas (Avg, Count), permitiendo construir vistas públicas de alta relevancia como el sistema de recomendaciones desarrollado."
    ]
    for con in conclusiones:
        p_c = doc.add_paragraph(con)
        p_c.style.font.name = "Arial"
        p_c.style.font.size = Pt(10.5)

    doc.save("Laboratorio5_Israel_Huaman.docx")
    print("Documento Word generado exitosamente: Laboratorio5_Israel_Huaman.docx")

if __name__ == "__main__":
    create_report()
